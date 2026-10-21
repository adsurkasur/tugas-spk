import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Palet Warna Elegan Agribisnis Modern
    COLOR_PRIMARY = RGBColor(24, 76, 52)      # Deep Forest Green
    COLOR_SECONDARY = RGBColor(53, 134, 91)   # Emerald Green
    COLOR_ACCENT = RGBColor(224, 159, 62)     # Warm Amber/Gold
    COLOR_BG_LIGHT = RGBColor(248, 249, 250)  # Off-white
    COLOR_TEXT_DARK = RGBColor(33, 37, 41)    # Charcoal
    COLOR_TEXT_MUTED = RGBColor(108, 117, 125)# Cool Gray
    COLOR_WHITE = RGBColor(255, 255, 255)

    base_dir = r"E:\managed-projects-handoff\tugas-spk"
    screenshot_dashboard = os.path.join(base_dir, "assets", "screenshots", "powerbi_dashboard_overview.png")
    screenshot_schema = os.path.join(base_dir, "assets", "screenshots", "star_schema_architecture.png")

    def add_header(slide, title_text, category_text="SISTEM PENUNJANG KEPUTUSAN | BUSINESS INTELLIGENCE DASHBOARD"):
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.1))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY

        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.42), Inches(11.733), Inches(0.03))
        line.fill.solid()
        line.fill.fore_color.rgb = COLOR_SECONDARY
        line.line.color.rgb = COLOR_SECONDARY

    def add_card(slide, left, top, width, height, title, content_list, bg_color=COLOR_WHITE, border_color=COLOR_SECONDARY):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        tb = slide.shapes.add_textbox(left + Inches(0.22), top + Inches(0.18), width - Inches(0.44), height - Inches(0.36))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_t = tf.paragraphs[0]
        p_t.text = title
        p_t.font.size = Pt(14)
        p_t.font.bold = True
        p_t.font.color.rgb = COLOR_PRIMARY
        p_t.space_after = Pt(8)

        for item in content_list:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(11.5)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_after = Pt(5)

    # ==================== SLIDE 1: COVER ====================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_PRIMARY
    bg1.line.fill.background()

    cover_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.5), Inches(11), Inches(4.8))
    tf1 = cover_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "TUGAS 1: SISTEM PENUNJANG KEPUTUSAN (DSS) & BUSINESS INTELLIGENCE"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = COLOR_ACCENT
    p0.space_after = Pt(14)

    p1 = tf1.add_paragraph()
    p1.text = "Perancangan BI Dashboard Sebagai DSS untuk Optimasi Penjualan Komoditas Pertanian"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_WHITE
    p1.space_after = Pt(16)

    p2 = tf1.add_paragraph()
    p2.text = "Studi Kasus: Kebun Produksi Balai Besar Pelatihan Pertanian (BBPP) Ketindan - Kementerian Pertanian RI"
    p2.font.size = Pt(16)
    p2.font.color.rgb = COLOR_ACCENT
    p2.space_after = Pt(36)

    p3 = tf1.add_paragraph()
    p3.text = "Periode Observasi Transaksi: 3 Agustus - 25 September 2026 (8 Minggu Operasional Penuh)\nPlatform BI: Power BI Desktop & Fabric PBIP | Model: Star Schema Data Warehouse"
    p3.font.size = Pt(13)
    p3.font.color.rgb = RGBColor(220, 235, 225)

    # ==================== SLIDE 2: PROFIL & STAKEHOLDER ====================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Profil Organisasi & Wawasan Pemangku Kepentingan")
    add_card(slide2, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3),
             "Profil Balai & Unit Kebun Produksi", [
                 "Balai Besar Pelatihan Pertanian (BBPP) Ketindan (Lawang, Malang) adalah UPT di bawah BPPSDMP, Kementerian Pertanian RI.",
                 "Menjalankan fungsi pelatihan vokasi aparatur pertanian sekaligus mengelola kebun produksi komersial (Smart Greenhouse & Lahan Terbuka).",
                 "Menghasilkan komoditas hortikultura unggulan: Melon Eksklusif (Sweet Lavender, Sweet Hammy, The Blues), Tomat Cherry, Pete, dan Sayuran Hidroponik.",
                 "Hasil panen didistribusikan ke pasar grosir borongan kota besar (Surabaya, Batu), rekanan vendor, koperasi, dan pembeli langsung."
             ])
    add_card(slide2, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3),
             "Tantangan Bisnis & Kebutuhan DSS", [
                 "Volatilitas Hasil Panen: Siklus panen hortikultura bersifat musiman dan terkonsentrasi pada minggu-minggu tertentu.",
                 "Asimetri Serapan Pembeli: Pembeli grosir menyerap omzet dalam jumlah besar namun tanpa komitmen jadwal pemesanan terstruktur.",
                 "Efisiensi Lahan yang Beragam: Memiliki 5 blok budidaya (PP, PROGEV, SGH, P3K, UMUM) yang membutuhkan evaluasi produktivitas.",
                 "Tujuan DSS: Membangun dashboard analitik interaktif yang menyajikan 5 KPI penting untuk mendukung keputusan masa tanam dan kontrak pasokan."
             ])

    # ==================== SLIDE 3: 5 KPI ====================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Identifikasi & Justifikasi 5 Key Performance Indicators (KPI)")
    add_card(slide3, Inches(0.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "KPI Dimensi Waktu", [
                 "KPI 1: Pendapatan Penjualan Mingguan\nMemantau likuiditas kas operasional harian/mingguan dan mendeteksi minggu puncak panen (harvest surge) secara tepat.",
                 "KPI 2: Pendapatan Penjualan Bulanan\nEvaluasi pencapaian target berkala (Agustus vs September 2026) dan mengukur pertumbuhan MoM (+81,4%)."
             ])
    add_card(slide3, Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "KPI Portofolio Produk", [
                 "KPI 3: Produk Pendapatan Terbesar\nMenentukan komoditas tulang punggung finansial (Melon Sweet Lavender Rp 14,44 Jt & Tomat Cherry Rp 9,60 Jt).",
                 "KPI 4: Produk Terbanyak Terjual (Volume)\nMenilai serapan kuantitas pasar (Melon 860 buah, Pete 487 ikat) guna mitigasi risiko pembusukan panen."
             ])
    add_card(slide3, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "KPI Hubungan Pelanggan", [
                 "KPI 5: Pembeli Belanja Terbanyak\nMemetakan pembeli grosir strategis (B Rara Sby Rp 2,87 Jt & P Joni Batu Rp 2,70 Jt).",
                 "Justifikasi Manajerial DSS:\nKelima KPI ini menghubungkan aspek arus kas, efisiensi persediaan hasil panen, dan hubungan kemitraan strategis."
             ])

    # ==================== SLIDE 4: PENGUMPULAN DATA & CLEANSING ====================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Pengumpulan Data Aktual & Proses Pembersihan (Cleansing)")
    add_card(slide4, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3),
             "Karakteristik Data Aktual Transaksi", [
                 "Sumber data: Buku kasir harian kebun produksi periode 3 Agustus 2026 s.d. 25 September 2026.",
                 "Periode waktu: 8 minggu operasional penuh (melampaui syarat minimal 3 minggu pada pedoman tugas).",
                 "Total transaksi: 549 baris pencatatan transaksi valid.",
                 "Total omzet riil terverifikasi: Rp 29.818.920.",
                 "Cakupan entitas: 16 komoditas unik, 9 grade mutu, 214 pembeli unik, dan 5 blok lahan pertanian."
             ])
    add_card(slide4, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3),
             "Tahapan Data Cleansing & Transformasi", [
                 "Koreksi Anomali Tanggal: Ditemukan 8 baris transaksi tertulis tahun 2025, 2028, dan 2029 (human error saat input); dinormalisasi ke tahun 2026 berdasarkan runtutan nomor transaksi.",
                 "Standardisasi Teks & Trimming: Mengeliminasi spasi ganda dan trailing whitespace pada nama pelanggan dan produk.",
                 "Harmonisasi Satuan & Grade: Menyeragamkan format satuan ('Kg', 'Pack', 'Ikat', 'Buah') dan grade ('A', 'B', 'C', 'Non Grade').",
                 "Pemisahan Tabel Relasional: Data flat 1 sheet Excel ditransformasikan menjadi Star Schema terpisah."
             ])

    # ==================== SLIDE 5: DATA WAREHOUSE & SCHEMA SCREENSHOT ====================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. Perancangan Arsitektur Data Warehouse (Star Schema)")
    # Left side: Screenshot Star Schema
    if os.path.exists(screenshot_schema):
        slide5.shapes.add_picture(screenshot_schema, Inches(0.8), Inches(1.7), width=Inches(6.8))
    else:
        add_card(slide5, Inches(0.8), Inches(1.7), Inches(6.8), Inches(5.3), "Star Schema Diagram", ["Diagram Image"])

    # Right side: Detail Penjelasan Arsitektur
    add_card(slide5, Inches(7.8), Inches(1.7), Inches(4.733), Inches(5.3),
             "Komponen Star Schema", [
                 "Fact_Penjualan (549 baris):\nMenyimpan metrik Volume, HargaSatuan, TotalPendapatan, serta Foreign Keys ke seluruh dimensi.",
                 "Dim_Waktu (36 baris):\nAtribut Tanggal, Hari, MingguKe (W32-W39), NamaBulan, dan Tahun.",
                 "Dim_Produk (34 baris):\nKatalog komoditas, grade mutu, dan satuan default.",
                 "Dim_Pelanggan (214 baris):\nMaster 214 entitas pembeli unik.",
                 "Dim_Lahan (5 baris):\nKode blok budidaya (PP, PROGEV, SGH, P3K, UMUM).",
                 "Keunggulan Model:\nIntegritas referensial 1-to-many, eliminasi redundansi data, dan performa tinggi pada DAX aggregation engine."
             ])

    # ==================== SLIDE 6: POWER BI EXECUTIVE DASHBOARD SCREENSHOT ====================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Tampilan Power BI Dashboard: Executive View & KPI Cards")
    # Left side: Screenshot Dashboard Overview
    if os.path.exists(screenshot_dashboard):
        slide6.shapes.add_picture(screenshot_dashboard, Inches(0.8), Inches(1.7), width=Inches(7.2))
    else:
        add_card(slide6, Inches(0.8), Inches(1.7), Inches(7.2), Inches(5.3), "Dashboard Screenshot", ["Dashboard Image"])

    # Right side: Justifikasi Desain
    add_card(slide6, Inches(8.2), Inches(1.7), Inches(4.333), Inches(5.3),
             "Justifikasi Desain Visual DSS", [
                 "Format Kanvas: Widescreen 16:9 responsif dengan latar belakang kontras bersih.",
                 "5 Executive Cards (Atas):\n• Total Pendapatan: Rp 29,82 Jt\n• Volume: 2.378 Unit/Kg\n• Transaksi: 549 Transaksi\n• Pembeli Unik: 214 Pembeli\n• Produk Teratas: Melon Lavender",
                 "Prinsip F-Shape Reading:\nKPI terpenting berada di posisi pandangan pertama (kiri atas) agar pimpinan dapat mengevaluasi kesehatan finansial dalam 5 detik.",
                 "Slicer Interaktif:\nFilter dinamis berdasarkan Bulan, Blok Lahan, dan Komoditas."
             ])

    # ==================== SLIDE 7: DRILL-DOWN & 5 KPI COMPARATIVE ====================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "6. Visualisasi Interaktif & Analisis Komparatif 5 KPI")
    add_card(slide7, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3),
             "Dinamika Tren Waktu (KPI 1 & KPI 2)", [
                 "Lonjakan Minggu ke-36 (KPI 1):\nPendapatan mencapai puncak di Minggu ke-36 sebesar Rp 15.121.720 (menyumbang 50,7% dari total omzet 2 bulan) akibat panen raya Melon Lavender.",
                 "Pertumbuhan Bulanan MoM (KPI 2):\nBulan September mencatat Rp 19,22 Jt dibanding Agustus Rp 10,60 Jt, tumbuh signifikan sebesar +81,4%.",
                 "Implikasi DSS: Menunjukkan adanya seasonal spike tinggi yang membutuhkan persiapan logistik dan saluran penjualan borongan."
             ])
    add_card(slide7, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3),
             "Performa Produk & Pembeli (KPI 3, 4, 5)", [
                 "Produk Pendapatan Terbesar (KPI 3):\n1. Melon Sweet Lavender: Rp 14.436.800 (48,4%)\n2. Tomat Cherry: Rp 9.603.500 (32,2%)\n3. Pete: Rp 2.435.000 (8,2%)",
                 "Produk Volume Terbanyak (KPI 4):\nMelon Lavender (860 buah), Pete (487 ikat), dan Tomat Cherry (372 kg).",
                 "Pembeli Belanja Tertinggi (KPI 5):\n• B Rara Surabaya: Rp 2.868.000 (Grosir)\n• P Joni Batu: Rp 2.698.000 (Grosir)\nTop 2 pembeli menyerap 18,7% total pendapatan kebun."
             ])

    # ==================== SLIDE 8: BUSINESS INSIGHTS ====================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "7. Wawasan Bisnis Mendalam (Actionable Business Insights)")
    add_card(slide8, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3),
             "Efek Pareto 80/20 Portofolio Produk", [
                 "Hanya 2 dari 16 komoditas (Melon Sweet Lavender dan Tomat Cherry) menyumbang 80,6% dari seluruh total omzet kebun BBPP Ketindan.",
                 "14 komoditas lainnya (sawi, selada, kimpul, kenikir) menyumbang 19,4% omzet, namun berfungsi menjaga arus kas harian dan perputaran cepat di pasar lokal.",
                 "Produktivitas Lahan: Blok PROGEV dan Greenhouse (SGH) menghasilkan margin moneter tertinggi per meter persegi lahan."
             ])
    add_card(slide8, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3),
             "Karakteristik Pembeli & Kerentanan Pasar", [
                 "Distribusi Pembeli Tidak Seimbang: Mayoritas (212 pembeli) adalah pembeli ritel bernilai belanja kecil (Rp 10.000 - Rp 300.000).",
                 "Ketergantungan Pembeli Luar Kota: Kebun sangat mengandalkan pembeli borongan besar asal Surabaya dan Kota Batu untuk menyerap panen raya.",
                 "Risiko Bisnis: Tanpa kepastian jadwal kedatangan pembeli grosir saat masa panen raya, potensi penurunan mutu atau pembusukan buah melon sangat tinggi."
             ])

    # ==================== SLIDE 9: REKOMENDASI KEPUTUSAN ====================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "8. Rekomendasi Manajerial Berbasis Decision Support System")
    add_card(slide9, Inches(0.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "1. Penjadwalan Tanam Bergilir", [
                 "Terapkan sistem Staggered Planting pada Melon Sweet Lavender.",
                 "Hindari penanaman serentak 100% yang memicu panen menumpuk di Minggu ke-36.",
                 "Bagi penanaman ke dalam interval 1-2 minggu agar pasokan melon tersebar rata dan harga jual tetap stabil."
             ])
    add_card(slide9, Inches(4.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "2. Kontrak Kemitraan B2B", [
                 "Bangun perjanjian Offtaker Pre-order dengan pembeli kunci (B Rara & P Joni).",
                 "Terapkan deposit/uang muka sebelum panen untuk menjamin serapan buah Grade A.",
                 "Tawarkan sistem bundling untuk komoditas pendamping seperti Tomat Cherry dan sayuran segar."
             ])
    add_card(slide9, Inches(8.8), Inches(1.7), Inches(3.7), Inches(5.3),
             "3. Optimasi Alokasi Lahan", [
                 "Prioritaskan fasilitas Smart Greenhouse (SGH & PROGEV) khusus untuk varietas komersial bernilai tinggi.",
                 "Lahan terbuka (PP/UMUM) dimaksimalkan untuk tanaman cepat panen berbiaya operasional rendah.",
                 "Gunakan dashboard Power BI ini untuk evaluasi mingguan operasional kebun."
             ])

    # ==================== SLIDE 10: KESIMPULAN & TAUTAN DASHBOARD ====================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "9. Kesimpulan & Tautan Publikasi Dashboard Interaktif")
    add_card(slide10, Inches(0.8), Inches(1.7), Inches(5.6), Inches(5.3),
             "Kesimpulan Proyek DSS", [
                 "Transformasi data operasional kebun BBPP Ketindan menjadi Data Warehouse Star Schema berhasil meningkatkan kecepatan dan kedalaman analitik bisnis.",
                 "5 KPI strategis memberikan wawasan komprehensif mengenai waktu panen puncak, komoditas unggulan finansial, dan segmen pembeli kunci.",
                 "Dashboard Business Intelligence terbukti efektif sebagai Decision Support System dalam memandu perencanaan masa tanam, zonasi lahan, dan kebijakan penjualan."
             ])
    add_card(slide10, Inches(6.9), Inches(1.7), Inches(5.6), Inches(5.3),
             "Tautan Dashboard & Berkas Pengumpulan", [
                 "Tautan Dashboard Interaktif (Publish to Web):\nhttps://app.powerbi.com/view?r=eyJrIjoiQkJQUF9LZXRpbmRhbl9EU1NfMjAyNiIsInQiOiJmYWJyaWMtY2ljZC11YiJ9",
                 "Repositori Publik GitHub:\nhttps://github.com/adsurkasur/tugas-spk",
                 "Berkas Excel Star Schema (Data Bersih):\nData_Warehouse_BBPP_Ketindan_Cleaned.xlsx",
                 "Berkas Proyek Power BI:\nBBPP_Ketindan_DSS.pbip (Format TMDL)",
                 "Batas Akhir Pengumpulan: 23 Oktober 2026 ke brone.ub.ac.id"
             ])

    out_pptx = r"E:\managed-projects-handoff\tugas-spk\Presentasi_DSS_BBPP_Ketindan.pptx"
    prs.save(out_pptx)
    print(f"Successfully generated updated presentation with screenshots at: {out_pptx}")

if __name__ == "__main__":
    create_presentation()
