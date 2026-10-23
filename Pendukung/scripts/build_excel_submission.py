"""Lengkapi Excel pengumpulan di root repo:
- 3_Data Warehouse: sheet Data_Mentah (salinan buku kasir asli), Ringkasan_KPI
  (jawaban 5 KPI), dan satuan Fact_Penjualan yang diseragamkan.
- 2_Data Bersih: satu tabel datar hasil pembersihan (fakta + dimensi digabung).
Aman dijalankan ulang. Diakhiri self-check terhadap Pendukung/data_csv/."""
import collections
import csv
import os

import openpyxl
from openpyxl.styles import Font, PatternFill

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # Pendukung/
ROOT = os.path.dirname(BASE)
PREFIX = "SPK_Ade Surya Ananda_2646000083"
RAW = os.path.join(ROOT, f"{PREFIX}_1_Buku Kasir Harian.xlsx")
CLEAN = os.path.join(ROOT, f"{PREFIX}_2_Data Bersih.xlsx")
OUT = os.path.join(ROOT, f"{PREFIX}_3_Data Warehouse Star Schema.xlsx")
SATUAN = {"kg": "Kg", "paket": "Pack"}  # paket & Pack berharga sama (Pete Rp5.000)
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="184C34")


def sheet_rows(ws):
    rows = list(ws.iter_rows(values_only=True))
    return [dict(zip(rows[0], r)) for r in rows[1:]]


def write_table(ws, row, title, header, data):
    ws.cell(row, 1, title).font = Font(bold=True, size=12)
    for j, h in enumerate(header, 1):
        c = ws.cell(row + 1, j, h)
        c.font, c.fill = HEAD, FILL
    for i, r in enumerate(data, row + 2):
        for j, v in enumerate(r, 1):
            c = ws.cell(i, j, v)
            if isinstance(v, float) or (j > 1 and isinstance(v, int)):
                c.number_format = "#,##0.##"
    return row + len(data) + 3


def main():
    wb = openpyxl.load_workbook(OUT)
    for name in ("Data_Mentah", "Ringkasan_KPI"):
        if name in wb.sheetnames:
            del wb[name]

    fact = wb["Fact_Penjualan"]
    col = [c.value for c in fact[1]].index("Satuan") + 1
    for (c,) in fact.iter_rows(min_row=2, min_col=col, max_col=col):
        c.value = SATUAN.get(c.value, c.value)

    raw = openpyxl.load_workbook(RAW, data_only=True).active
    ws = wb.create_sheet("Data_Mentah", 0)
    for r in raw.iter_rows(values_only=True):
        ws.append(list(r[:10]))
    for c in ws[3]:
        c.font, c.fill = HEAD, FILL

    f = sheet_rows(fact)
    waktu = {r["DateKey"]: r for r in sheet_rows(wb["Dim_Waktu"])}
    produk = {r["ProductKey"]: r for r in sheet_rows(wb["Dim_Produk"])}
    pembeli = {r["CustomerKey"]: r["NamaPembeli"] for r in sheet_rows(wb["Dim_Pelanggan"])}

    def agg(key, val="TotalPendapatan"):
        d = collections.Counter()
        for r in f:
            d[key(r)] += r[val]
        return d

    total = sum(r["TotalPendapatan"] for r in f)
    minggu = sorted(agg(lambda r: (waktu[r["DateKey"]]["MingguKe"], waktu[r["DateKey"]]["LabelMinggu"])).items())
    bulan = sorted(agg(lambda r: (waktu[r["DateKey"]]["BulanAngka"], waktu[r["DateKey"]]["NamaBulan"])).items())
    prod_rp = agg(lambda r: produk[r["ProductKey"]]["Komoditas"]).most_common()
    prod_vol = agg(lambda r: (produk[r["ProductKey"]]["Komoditas"], r["Satuan"]), "Volume").most_common()
    cust = agg(lambda r: pembeli[r["CustomerKey"]]).most_common(10)

    ws = wb.create_sheet("Ringkasan_KPI", 1)
    ws["A1"] = "Ringkasan 5 KPI - Kebun Produksi BBPP Ketindan (3 Agustus - 25 September 2026)"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Total pendapatan Rp {total:,.0f} dari {len(f)} transaksi; dihitung dari sheet Fact_Penjualan."
    pct = lambda v: round(v / total * 100, 1)
    row = write_table(ws, 4, "KPI 1 - Pendapatan per Minggu", ["Minggu", "Pendapatan (Rp)", "% Omzet"],
                      [(k[1], v, pct(v)) for k, v in minggu])
    row = write_table(ws, row, "KPI 2 - Pendapatan per Bulan", ["Bulan", "Pendapatan (Rp)", "% Omzet"],
                      [(k[1], v, pct(v)) for k, v in bulan])
    row = write_table(ws, row, "KPI 3 - Pendapatan per Produk", ["Komoditas", "Pendapatan (Rp)", "% Omzet"],
                      [(k, v, pct(v)) for k, v in prod_rp])
    row = write_table(ws, row, "KPI 4 - Volume Terjual per Produk", ["Komoditas", "Satuan", "Volume"],
                      [(k[0], k[1], round(v, 2)) for k, v in prod_vol])
    write_table(ws, row, "KPI 5 - 10 Pembeli dengan Belanja Terbanyak", ["Pembeli", "Total Belanja (Rp)", "% Omzet"],
                [(k, v, pct(v)) for k, v in cust])
    for letter, w in zip("ABC", (38, 20, 12)):
        ws.column_dimensions[letter].width = w

    wb.save(OUT)

    lahan = {r["LahanKey"]: r["KodeLahan"] for r in sheet_rows(wb["Dim_Lahan"])}
    clean = openpyxl.Workbook()
    ws = clean.active
    ws.title = "Data_Bersih"
    header = ["No", "Tanggal", "Hari", "Minggu", "Bulan", "Komoditas", "Grade", "Pembeli",
              "Volume", "Satuan", "Harga Satuan (Rp)", "Total (Rp)", "Lahan"]
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in sorted(f, key=lambda r: r["TransactionID"]):
        t, p = waktu[r["DateKey"]], produk[r["ProductKey"]]
        ws.append([r["TransactionID"], t["Tanggal"], t["Hari"], t["LabelMinggu"], t["NamaBulan"],
                   p["Komoditas"], p["Grade"], pembeli[r["CustomerKey"]], r["Volume"], r["Satuan"],
                   r["HargaSatuan"], r["TotalPendapatan"], lahan[r["LahanKey"]]])
    for row in ws.iter_rows(min_row=2):
        row[1].number_format = "dd/mm/yyyy"
        for c in row[10:12]:
            c.number_format = "#,##0"
    for letter, w in zip("ABCDEFGHIJKLM", (6, 12, 9, 11, 11, 22, 11, 34, 9, 8, 16, 14, 9)):
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    clean.save(CLEAN)

    # self-check: Excel harus identik dengan data_csv yang dipakai Power BI
    with open(os.path.join(BASE, "data_csv", "Fact_Penjualan.csv"), encoding="utf-8-sig") as fh:
        csv_rows = list(csv.DictReader(fh))
    assert len(csv_rows) == len(f) == 549
    assert abs(sum(float(r["TotalPendapatan"]) for r in csv_rows) - total) < 1 and total == 29818920
    assert {r["Satuan"] for r in csv_rows} == {r["Satuan"] for r in f} == {"Kg", "Pack", "Ikat", "Buah"}
    assert prod_rp[0][0] == "Melon Sweet Lavender" and cust[0][0] == "B Rara Surabaya"
    flat = list(openpyxl.load_workbook(CLEAN).active.iter_rows(min_row=2, values_only=True))
    assert len(flat) == 549 and sum(r[11] for r in flat) == total
    assert all(r[1].year == 2026 for r in flat), "masih ada tanggal di luar 2026"
    print("OK", wb.sheetnames, "+", os.path.basename(CLEAN))


if __name__ == "__main__":
    main()
